"""Экспорт заказов в формат .docx.

Используется библиотека ``python-docx`` (входит в технические
требования ТЗ). Документ оформлен в виде таблицы со столбцами:
№, дата, клиент, телефон, статус, сумма, состав. Сверху размещён
заголовок с указанием периода.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, timedelta
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt

from bot.database import Database, Order

logger = logging.getLogger(__name__)


def export_orders_to_docx(
    db: Database,
    days: int = 7,
    statuses: Optional[List[str]] = None,
) -> tuple[bytes, str]:
    """Сформировать .docx с заказами за указанный период.

    :param db: экземпляр :class:`Database`.
    :param days: количество последних дней (0 — за всё время).
    :param statuses: список фильтруемых статусов или ``None``.
    :return: tuple ``(содержимое файла, имя файла)``.
    """

    date_from: Optional[str] = None
    if days > 0:
        date_from = (datetime.now().date()
                     - timedelta(days=days - 1)).isoformat()
    orders: List[Order] = db.list_orders(
        statuses=statuses, limit=10000, offset=0,
        date_from=date_from,
    )

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Mm(20)
    section.right_margin = Mm(15)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Отчёт по заказам мебельного магазина")
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)
    run.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    period_text = ("за всё время" if days == 0
                   else f"за последние {days} дн.")
    sub = subtitle.add_run(
        f"Период: {period_text}.  "
        f"Дата формирования: "
        f"{datetime.now().strftime('%Y-%m-%d %H:%M')}."
    )
    sub.font.name = "Times New Roman"
    sub.font.size = Pt(12)
    sub.italic = True

    doc.add_paragraph()

    if not orders:
        p = doc.add_paragraph("За указанный период заказов нет.")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        headers = ["№", "Дата", "Клиент", "Телефон",
                   "Статус", "Сумма, ₽", "Состав"]
        table = doc.add_table(rows=1 + len(orders), cols=len(headers))
        table.style = "Table Grid"

        for j, h in enumerate(headers):
            cell = table.rows[0].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(h)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for i, o in enumerate(orders, start=1):
            row = table.rows[i].cells
            values = [
                str(o.id), o.created_at, o.customer_name,
                o.customer_phone, o.status_label,
                f"{o.total:,.2f}".replace(",", " "),
                o.items,
            ]
            for j, v in enumerate(values):
                row[j].text = ""
                r = row[j].paragraphs[0].add_run(v)
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                row[j].paragraphs[0].paragraph_format.line_spacing = 1.0

        doc.add_paragraph()
        summary = doc.add_paragraph()
        sr = summary.add_run(
            f"Итого заказов: {len(orders)}.  "
            f"Общая сумма: "
            f"{sum(o.total for o in orders):,.2f} ₽".replace(",", " ")
        )
        sr.font.name = "Times New Roman"
        sr.font.size = Pt(12)
        sr.bold = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    suffix = "all" if days == 0 else f"last{days}d"
    filename = (f"orders_{suffix}_"
                f"{datetime.now().strftime('%Y%m%d_%H%M')}.docx")
    return buf.read(), filename
