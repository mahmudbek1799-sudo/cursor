from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


OUT_DIR = Path("docs")
DOCX_OUT = OUT_DIR / "telegram_appendix_replacement.docx"
MD_OUT = OUT_DIR / "telegram_appendix_replacement.md"


APPENDIX_TEXT = """ПРИЛОЖЕНИЕ 1
(обязательное)
Фрагменты программного кода

В приложении приведены только основные фрагменты программного кода, которые отражают ключевые элементы реализации проекта: настройки приложения, работу с базой данных, получение данных с веб-страницы и формирование аналитического графика. Полные служебные обработчики интерфейса Telegram-бота в приложение не включены, так как они перегружают пояснительную записку и не относятся к минимальному набору, указанному в техническом задании.

1. Файл config.py

В файле config.py хранится класс Settings, который отвечает за чтение параметров окружения. Отдельно указывается параметр proxy_url, используемый при необходимости подключения через прокси-сервер.

```python
from pydantic_settings import BaseSettings, SettingsConfigDict


class Settings(BaseSettings):
    proxy_url: str | None = None

    model_config = SettingsConfigDict(
        env_file=".env",
        env_file_encoding="utf-8",
        extra="ignore",
    )


settings = Settings()
```

2. Файл database.py

В файле database.py приведены только функции init_schema, upsert_order и apply_discount. Они демонстрируют создание структуры базы данных, добавление или обновление заказа и применение скидки к заказу.

```python
import sqlite3
from pathlib import Path


DB_PATH = Path("data/app.db")


def get_connection() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    return sqlite3.connect(DB_PATH)


def init_schema() -> None:
    with get_connection() as conn:
        conn.execute(
            \"\"\"
            CREATE TABLE IF NOT EXISTS orders (
                order_id INTEGER PRIMARY KEY,
                user_id INTEGER NOT NULL,
                title TEXT NOT NULL,
                total REAL NOT NULL,
                discount REAL NOT NULL DEFAULT 0,
                updated_at TEXT DEFAULT CURRENT_TIMESTAMP
            )
            \"\"\"
        )
        conn.commit()


def upsert_order(order_id: int, user_id: int, title: str, total: float) -> None:
    with get_connection() as conn:
        conn.execute(
            \"\"\"
            INSERT INTO orders (order_id, user_id, title, total)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(order_id) DO UPDATE SET
                user_id = excluded.user_id,
                title = excluded.title,
                total = excluded.total,
                updated_at = CURRENT_TIMESTAMP
            \"\"\",
            (order_id, user_id, title, total),
        )
        conn.commit()


def apply_discount(order_id: int, percent: float) -> float:
    if not 0 <= percent <= 100:
        raise ValueError("Размер скидки должен быть от 0 до 100 процентов")

    with get_connection() as conn:
        row = conn.execute(
            "SELECT total FROM orders WHERE order_id = ?",
            (order_id,),
        ).fetchone()
        if row is None:
            raise ValueError("Заказ не найден")

        total = float(row[0])
        discount = round(total * percent / 100, 2)
        conn.execute(
            "UPDATE orders SET discount = ?, updated_at = CURRENT_TIMESTAMP WHERE order_id = ?",
            (discount, order_id),
        )
        conn.commit()
        return round(total - discount, 2)
```

3. Файл parser.py

В файле parser.py оставлен только фрагмент, показывающий разбор HTML-страницы с использованием BeautifulSoup. Такой фрагмент отражает принцип извлечения данных из разметки без включения лишней служебной логики.

```python
from bs4 import BeautifulSoup


def parse_orders_page(html: str) -> list[dict[str, str]]:
    soup = BeautifulSoup(html, "html.parser")
    result: list[dict[str, str]] = []

    for row in soup.select("table.orders tbody tr"):
        cells = [cell.get_text(strip=True) for cell in row.select("td")]
        if len(cells) < 4:
            continue

        result.append(
            {
                "order_id": cells[0],
                "client": cells[1],
                "title": cells[2],
                "total": cells[3].replace(" ", "").replace(",", "."),
            }
        )

    return result
```

4. Файл analytics.py

В файле analytics.py приведен только фрагмент генерации графика средствами matplotlib. Он показывает, как на основе агрегированных данных формируется изображение для дальнейшей отправки пользователю или сохранения в отчете.

```python
from pathlib import Path

import matplotlib.pyplot as plt


def build_orders_chart(data: dict[str, float], output_path: str | Path) -> Path:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    labels = list(data.keys())
    values = list(data.values())

    fig, ax = plt.subplots(figsize=(8, 4.5))
    ax.bar(labels, values, color="#4C78A8")
    ax.set_title("Сумма заказов по категориям")
    ax.set_xlabel("Категория")
    ax.set_ylabel("Сумма, руб.")
    ax.grid(axis="y", linestyle="--", alpha=0.4)

    fig.tight_layout()
    fig.savefig(output, dpi=150)
    plt.close(fig)

    return output
```
"""


def set_run_font(run, *, size=14, name="Times New Roman", bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def format_paragraph(paragraph, *, first_line=True, spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph.alignment = align
    paragraph.paragraph_format.left_indent = Cm(0)
    paragraph.paragraph_format.right_indent = Cm(0)
    paragraph.paragraph_format.first_line_indent = Cm(1.25) if first_line else Cm(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = spacing
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:widowControl")) is None:
        p_pr.append(OxmlElement("w:widowControl"))


def add_paragraph(doc: Document, text: str, *, bold=False, first_line=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, first_line=first_line, align=align)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    return paragraph


def add_code_line(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    format_paragraph(paragraph, first_line=False, spacing=1.0, align=WD_ALIGN_PARAGRAPH.LEFT)
    paragraph.paragraph_format.left_indent = Cm(1.0)
    run = paragraph.add_run(text)
    set_run_font(run, size=10, name="Courier New")


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(30)
    section.right_margin = Mm(15)

    in_code = False
    for raw_line in APPENDIX_TEXT.splitlines():
        line = raw_line.rstrip()
        if line == "```python":
            in_code = True
            continue
        if line == "```":
            in_code = False
            doc.add_paragraph()
            continue
        if in_code:
            add_code_line(doc, line)
            continue
        if not line:
            doc.add_paragraph()
            continue
        if line == "ПРИЛОЖЕНИЕ 1":
            add_paragraph(doc, line, bold=True, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line == "(обязательное)":
            add_paragraph(doc, line, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line == "Фрагменты программного кода":
            add_paragraph(doc, line, bold=True, first_line=False, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif line[0].isdigit() and ". Файл " in line:
            add_paragraph(doc, line, bold=True, first_line=False, align=WD_ALIGN_PARAGRAPH.LEFT)
        else:
            add_paragraph(doc, line)

    doc.core_properties.title = "Замена приложения к пояснительной записке"
    doc.core_properties.subject = "Сокращенное приложение с фрагментами кода"
    doc.save(DOCX_OUT)


def main():
    OUT_DIR.mkdir(exist_ok=True)
    MD_OUT.write_text(APPENDIX_TEXT, encoding="utf-8")
    build_docx()
    print(f"Created {DOCX_OUT}")
    print(f"Created {MD_OUT}")


if __name__ == "__main__":
    main()
